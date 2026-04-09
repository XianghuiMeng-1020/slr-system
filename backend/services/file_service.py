from __future__ import annotations
import csv
import io
import json
import os
import uuid
import zipfile

import openpyxl


def parse_coding_scheme(file_path: str, filename: str) -> list[dict]:
    """Parse a coding scheme file (CSV, XLSX, or JSON) into a list of dicts."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".json":
        return _parse_json(file_path)
    elif ext == ".csv":
        return _parse_csv(file_path)
    elif ext in (".xlsx", ".xls"):
        return _parse_xlsx(file_path)
    else:
        raise ValueError(f"Unsupported coding scheme format: {ext}")


def _parse_json(path: str) -> list[dict]:
    with open(path, "r", encoding="utf-8-sig") as f:
        data = json.load(f)
    return _normalize_json_data(data)


def _normalize_json_data(data: object) -> list[dict]:
    """Accept array of objects, or an object containing an array."""
    if isinstance(data, list):
        items = [x for x in data if isinstance(x, dict)]
        if items:
            return _normalize_items(items)
        str_items = [x for x in data if isinstance(x, str)]
        if str_items:
            return _normalize_items(
                [{"code": f"C{i+1}", "description": s} for i, s in enumerate(str_items)]
            )
        raise ValueError("JSON array contains no valid objects or strings")
    if isinstance(data, dict):
        for key in ("items", "codes", "scheme", "coding_scheme", "data"):
            if key in data and isinstance(data[key], list):
                return _normalize_json_data(data[key])
        return _normalize_items([data])
    raise ValueError("JSON must be an array or an object containing an array")


def _parse_csv(path: str) -> list[dict]:
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
    return _normalize_items(rows)


def _parse_xlsx(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip().lower() if h else "" for h in next(rows_iter)]
    items = []
    for row in rows_iter:
        item = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                item[headers[i]] = str(val).strip() if val is not None else ""
        if any(item.values()):
            items.append(item)
    wb.close()
    return _normalize_items(items)


def _normalize_items(items: list[dict]) -> list[dict]:
    """Normalize field names to: id, code, description, category."""
    result = []
    for i, item in enumerate(items):
        lower_item = {k.lower().strip(): v for k, v in item.items()}
        code = (
            lower_item.get("code")
            or lower_item.get("id")
            or lower_item.get("code_id")
            or lower_item.get("name")
            or f"C{i+1}"
        )
        description = (
            lower_item.get("description")
            or lower_item.get("desc")
            or lower_item.get("label")
            or lower_item.get("name")
            or code
        )
        category = (
            lower_item.get("category")
            or lower_item.get("group")
            or lower_item.get("theme")
            or ""
        )
        result.append({
            "id": str(uuid.uuid4())[:8],
            "code": str(code).strip(),
            "description": str(description).strip(),
            "category": str(category).strip() if category else None,
        })
    return result


def parse_coding_scheme_text(text: str) -> list[dict]:
    """Parse pasted text into coding scheme items. Supports JSON, CSV, or simple lines."""
    text = text.strip()
    if not text:
        raise ValueError("Empty text")

    if text.startswith("[") or text.startswith("{"):
        try:
            data = json.loads(text)
            return _normalize_json_data(data)
        except json.JSONDecodeError as e:
            raise ValueError(f"Invalid JSON: {e}")

    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        raise ValueError("No content found")

    if "," in lines[0] and len(lines[0].split(",")) >= 2:
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        if rows and any(rows[0].values()):
            return _normalize_items(rows)

    items = []
    for i, line in enumerate(lines):
        if ":" in line:
            parts = line.split(":", 1)
            code = parts[0].strip()
            desc = parts[1].strip()
        elif "\t" in line:
            parts = line.split("\t", 1)
            code = parts[0].strip()
            desc = parts[1].strip() if len(parts) > 1 else code
        else:
            code = f"C{i+1}"
            desc = line
        items.append({"code": code, "description": desc})

    return _normalize_items(items)


def extract_zip(zip_path: str, dest_dir: str) -> list[str]:
    """Extract PDFs from a ZIP archive, returns list of extracted file paths."""
    extracted = []
    with zipfile.ZipFile(zip_path, "r") as zf:
        for name in zf.namelist():
            if name.lower().endswith(".pdf") and not name.startswith("__MACOSX"):
                safe_name = os.path.basename(name)
                if not safe_name:
                    continue
                dest_path = os.path.join(dest_dir, f"{uuid.uuid4().hex[:8]}_{safe_name}")
                with zf.open(name) as src, open(dest_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append(dest_path)
    return extracted


def extract_docx_text(path: str) -> str:
    """Extract plain text from a Word .docx file."""
    try:
        from docx import Document as DocxDocument
    except ImportError as e:
        raise ValueError("python-docx is required for Word import") from e
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_html_text(html: str) -> str:
    """Strip HTML to plain text."""
    try:
        from bs4 import BeautifulSoup
    except ImportError as e:
        raise ValueError("beautifulsoup4 is required for HTML import") from e
    return BeautifulSoup(html, "html.parser").get_text("\n")
